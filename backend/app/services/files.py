from __future__ import annotations

import io
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document

ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class DocumentReadError(ValueError):
    pass


def extract_document(filename: str, content: bytes, max_chars: int = 200_000) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise DocumentReadError("Supported files are TXT, Markdown, PDF, and DOCX.")
    if not content:
        raise DocumentReadError("The uploaded file is empty.")
    if suffix in {".txt", ".md"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentReadError("Text files must use UTF-8 encoding.") from exc
    elif suffix == ".pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.is_encrypted:
                    raise DocumentReadError("Password-protected PDFs are not supported.")
                if len(document) > 5_000:
                    raise DocumentReadError("The PDF exceeds the 5,000-page extraction limit.")
                pieces = []
                extracted_chars = 0
                for page_number, page in enumerate(document, start=1):
                    text_page = page.get_text("text")
                    if text_page.strip():
                        piece = f"\n[Page {page_number}]\n{text_page}"
                        pieces.append(piece)
                        extracted_chars += len(piece)
                    if extracted_chars > max_chars:
                        break
                text = "\n".join(pieces)
        except DocumentReadError:
            raise
        except Exception as exc:
            raise DocumentReadError("Could not read this PDF. It may be damaged or unsupported.") from exc
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                infos = archive.infolist()
                if len(infos) > 5000 or sum(info.file_size for info in infos) > 50 * 1024 * 1024:
                    raise DocumentReadError("DOCX archive expands beyond the safe extraction limit.")
                if "word/document.xml" not in archive.namelist():
                    raise DocumentReadError("The DOCX file is missing its main document part.")
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        except DocumentReadError:
            raise
        except Exception as exc:
            raise DocumentReadError("Could not read this DOCX file. It may be damaged or unsupported.") from exc
    text = text.replace("\x00", "")
    if not text.strip():
        raise DocumentReadError("No readable text was found. Scanned PDFs need OCR before upload.")
    if len(text) > max_chars:
        raise DocumentReadError(f"Extracted text exceeds the {max_chars:,}-character limit; no content was silently truncated.")
    return re.sub(r"[ \t]+\n", "\n", text).strip()
